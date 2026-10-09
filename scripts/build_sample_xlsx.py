"""Build a formatted Excel version of the illustrative dataset of every scenario.

Reads data/illustrative/<ID>/*.csv (written by build_illustrative.py) and writes
frontend/demo/sample/<ID>_sample_dataset.xlsx and a copy in data/illustrative/<ID>/. Values are copied as computed by the
engine (no arithmetic here, PROJECT_RULES.md rule 1); this script only adds labels and formatting.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
IDS = {"S1": "Normal quarter", "S2": "Large customer slows down", "S3": "Construction sector slowdown",
       "S4": "One-week payment spike", "S5": "Seasonal slowdowns", "S6": "New customer, slow payer",
       "S7": "Partial payments", "S8": "Late month-end close", "H1": "Unseen test: customer drift",
       "H2": "Unseen test: sector slowdown"}

SEG = {"general_contractors": "General contractors", "specialised_installers": "Specialised installers",
       "property_developers": "Property developers", "public_entities": "Public entities", "retail_diy": "Retail / DIY"}

# header -> (label, kind, width)
COLS = {
    "week": ("Week", "int", 7), "week_start": ("Week starting", "date", 13),
    "opening_eur": ("Opening (€)", "eur", 13), "receipts_eur": ("Receipts (€)", "eur", 13),
    "pay_supplier_invoices_received_eur": ("Supplier invoices already received (€)", "eur", 16),
    "pay_purchases_for_new_sales_eur": ("Purchases for new sales (€)", "eur", 14),
    "pay_known_commitments_eur": ("Known large bills (€)", "eur", 13),
    "pay_payroll_rent_overheads_tax_eur": ("Payroll, rent, overheads, tax (€)", "eur", 15),
    "pay_vat_eur": ("VAT (€)", "eur", 12), "payments_total_eur": ("Total payments (€)", "eur", 14),
    "closing_expected_eur": ("Closing, expected (€)", "eur", 14), "closing_p10_eur": ("Low, P10 (€)", "eur", 13),
    "closing_p90_eur": ("High, P90 (€)", "eur", 13), "prob_below_safety_level": ("Chance below safety level", "pct", 12),
    "closing_if_all_paid_on_time_eur": ("If all paid on time (€)", "eur", 14),
    "closing_stress_largest_customer_plus30d_eur": ("Stress: largest customer +30 days (€)", "eur", 16),
    "safety_level_eur": ("Safety level (€)", "eur", 12), "closing_actual_eur": ("Actual, revealed later (€)", "eur", 14),
    "customer": ("Customer", "text", 11), "segment": ("Segment", "seg", 22),
    "owed_eur": ("Owed (€)", "eur", 12), "overdue_eur": ("Overdue (€)", "eur", 12),
    "usual_delay_days": ("Usual delay (days)", "num", 10), "recent_delay_days": ("Recent delay, 90 days (days)", "num", 11),
    "note": ("What it means", "wrap", 60), "invoice": ("Invoice", "text", 19), "due_date": ("Due date", "date", 12),
    "outstanding_eur": ("Amount open (€)", "eur", 13), "days_overdue": ("Days overdue", "int", 9),
    "predicted_payment": ("Predicted payment", "date", 12), "earliest_likely": ("Earliest likely", "date", 12),
    "latest_likely": ("Latest likely", "date", 12), "remind_by": ("Remind by", "date", 12),
}

SHEETS = [
    ("13-week plan", "weekly_forecast.csv",
     "13-week cash plan from 28 Sep 2026. Red rows: expected closing balance below the safety level."),
    ("Top customers", "top_customers.csv",
     "The 10 customers who owe the most, with their usual and recent payment delay."),
    ("Open invoices", "open_invoices.csv",
     "Every open invoice of those 10 customers, with predicted payment date. Red rows: already overdue."),
]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="0F2A3D")
RED_FILL = PatternFill("solid", fgColor="FBE3E1")
THIN = Border(bottom=Side(style="thin", color="D9D6CF"))


def convert(v, kind):
    if v == "":
        return None
    if kind == "eur":
        return float(v)
    if kind == "int":
        return int(float(v))
    if kind in ("num", "pct"):
        return float(v)
    if kind == "date":
        from datetime import date
        return date.fromisoformat(v)
    if kind == "seg":
        return SEG.get(v, v)
    return v


def add_sheet(wb, src, title, fname, desc):
    ws = wb.create_sheet(title)
    rows = list(csv.reader(open(src / fname, newline="", encoding="utf-8")))
    head, body = rows[0], rows[1:]
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = desc
    ws["A2"].font = Font(name=FONT, italic=True, size=10, color="5B6770")
    hr = 4
    for j, h in enumerate(head, 1):
        label, kind, width = COLS.get(h, (h, "text", 14))
        c = ws.cell(hr, j, label)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if kind in ("text", "seg", "date", "wrap") else "right")
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[hr].height = 42
    i_close = head.index("closing_expected_eur") if "closing_expected_eur" in head else None
    i_safe = head.index("safety_level_eur") if "safety_level_eur" in head else None
    i_od = head.index("days_overdue") if "days_overdue" in head else None
    for i, r in enumerate(body, hr + 1):
        bad = (i_close is not None and float(r[i_close]) < float(r[i_safe])) or (i_od is not None and float(r[i_od]) > 0)
        for j, (h, v) in enumerate(zip(head, r), 1):
            kind = COLS.get(h, (h, "text", 14))[1]
            c = ws.cell(i, j, convert(v, kind))
            c.font = Font(name=FONT, size=10)
            c.border = THIN
            if kind == "eur":
                c.number_format = '#,##0;-#,##0;"-"'
            elif kind == "pct":
                c.number_format = '0%;-0%;"-"'
            elif kind == "num":
                c.number_format = "0"
            elif kind == "date":
                c.number_format = "dd mmm yyyy"
                c.alignment = Alignment(horizontal="left")
            elif kind == "wrap":
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if bad:
                c.fill = RED_FILL
    ws.freeze_panes = ws.cell(hr + 1, 2)
    ws.sheet_view.showGridLines = False


def build(sid):
    src = ROOT / "data" / "illustrative" / sid
    out = ROOT / "frontend" / "demo" / "sample" / f"{sid}_sample_dataset.xlsx"
    copy = src / f"{sid}_sample_dataset.xlsx"  # same workbook next to the CSV files and README
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    lines = [
        (f"Headroom Monitor · illustrative dataset (scenario {sid} · {IDS[sid]})", Font(name=FONT, bold=True, size=14)),
        ("Fictional company DemoCo Building Supplies S.à r.l. · cash position as of 28 Sep 2026", Font(name=FONT, size=10, color="5B6770")),
        ("", None),
        ("What is in this file", Font(name=FONT, bold=True, size=11)),
        ("13-week plan: opening balance + receipts − payments = closing balance, week by week, with the likely range and the safety level.", None),
        ("Top customers: the 10 customers who owe the most and how late they usually pay.", None),
        ("Open invoices: every open invoice of those customers, with the predicted payment date and when to send a reminder.", None),
        ("", None),
        ("How to read it", Font(name=FONT, bold=True, size=11)),
        ("Red rows = a problem: closing balance below the safety level, or an invoice already overdue.", None),
        ("Expected = middle of 1,000 simulations. Low (P10) / High (P90) = likely range.", None),
        ("Safety level = 4 weeks of fixed costs. Actual = the hidden future, revealed only after the forecast was made.", None),
        ("All figures are computed by the forecasting engine and rounded to the euro, so a total can differ from the sum of its parts by €1.", None),
        ("", None),
        ("All data is synthetic. Customer codes (CUST_###) are pseudonyms.", Font(name=FONT, italic=True, size=10, color="5B6770")),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = f or Font(name=FONT, size=10)
    ws.column_dimensions["A"].width = 120
    ws.sheet_view.showGridLines = False
    for title, fname, desc in SHEETS:
        add_sheet(wb, src, title, fname, desc)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    wb.save(copy)
    print("wrote", out.relative_to(ROOT), "and", copy.relative_to(ROOT))


def main():
    for sid in IDS:
        build(sid)


if __name__ == "__main__":
    main()
